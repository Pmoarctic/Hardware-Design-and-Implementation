from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / 'stopwatch-interrupt-th.docx'
img = Image.new('RGB', (1500, 1230), 'white')
d = ImageDraw.Draw(img)
INK, LINE = '#172536', '#4A6175'

def font(size, bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/' + ('arialbd.ttf' if bold else 'arial.ttf'), size)

def box(rect, lines, fill='#F1F5F9', size=32):
    d.rounded_rectangle(rect, radius=20, fill=fill, outline=LINE, width=3)
    x0,y0,x1,y1=rect
    height=len(lines)*40
    y=(y0+y1-height)/2
    for i,line in enumerate(lines):
        d.text(((x0+x1)/2,y+i*40),line,font=font(size,i==0),fill=INK,anchor='mt')

def arrow(points, color=LINE):
    d.line(points,fill=color,width=5,joint='curve')
    x,y=points[-1]; px,py=points[-2]
    if y>py: tri=[(x,y),(x-11,y-18),(x+11,y-18)]
    elif y<py: tri=[(x,y),(x-11,y+18),(x+11,y+18)]
    elif x>px: tri=[(x,y),(x-18,y-11),(x-18,y+11)]
    else: tri=[(x,y),(x+18,y-11),(x+18,y+11)]
    d.polygon(tri,fill=color)

box((470,20,1030,135),['Main program','Wait for an interrupt'])
box((140,185,640,335),['Timer tick','Every 1,000 CPU cycles','Request INT0'],'#EAF3FC',30)
box((860,185,1360,335),['Press LAP','Request INT2'],'#E8F6EF')
box((470,390,1030,505),['CPU accepts interrupt','Save return address'])
arrow([(750,135),(750,390)])
arrow([(390,335),(390,445),(470,445)])
arrow([(1110,335),(1110,445),(1030,445)])
box((470,555,1030,655),['Handler saves registers'])
arrow([(750,505),(750,555)])
box((140,735,640,885),['INT0  Timer handler','If running:','elapsed = elapsed + 1'],'#EAF3FC',32)
box((860,735,1360,885),['INT2  Lap handler','If running and space remains:','save time and lap duration'],'#E8F6EF',30)
arrow([(750,655),(750,695),(390,695),(390,735)])
arrow([(750,695),(1110,695),(1110,735)])
box((470,945,1030,1035),['Restore registers'])
arrow([(390,885),(390,990),(470,990)])
arrow([(1110,885),(1110,990),(1030,990)])
box((470,1090,1030,1200),['RETI','Return to main program'])
arrow([(750,1035),(750,1090)])
arrow([(470,1145),(55,1145),(55,78),(470,78)])
img.save(HERE/'interrupt-diagram.png')

doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.27)
section.page_height=Inches(11.69)
section.top_margin=Inches(.60)
section.bottom_margin=Inches(.60)
section.left_margin=Inches(.72)
section.right_margin=Inches(.72)

for name in ['Normal','Title','Heading 1']:
    style=doc.styles[name]
    style.font.name='Tahoma'
    style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(11)
    rpr=style.element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:
        rf=OxmlElement('w:rFonts');rpr.append(rf)
    for key in ['ascii','hAnsi','eastAsia','cs']:
        rf.set(qn('w:'+key),'Tahoma')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'th-TH');rpr.append(lang)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.line_spacing=1.13
doc.styles['Title'].font.size=Pt(19)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(8)

doc.add_paragraph('การทำงานของ Interrupt ในนาฬิกาจับเวลา',style='Title')
p=doc.add_paragraph()
p.add_run('Interrupt คือการขัดจังหวะชั่วคราว ').bold=True
p.add_run('CPU พักงานหลักเพื่อจัดการเหตุการณ์ที่เข้ามา เมื่อทำเสร็จก็กลับไปทำงานต่อจากจุดเดิม เช่น รับสัญญาณนับเวลา หรือบันทึกเวลาเมื่อกดปุ่ม Lap')

p=doc.add_paragraph()
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before=Pt(3)
p.paragraph_format.space_after=Pt(8)
r=p.add_run(); shape=r.add_picture(str(HERE/'interrupt-diagram.png'),width=Inches(6.64))
shape._inline.docPr.set('descr','แผนภาพจากงานหลักไปยังตัวจัดการ INT0 เพื่อนับเวลา หรือ INT2 เพื่อบันทึก Lap แล้วคืนค่ารีจิสเตอร์และกลับด้วย RETI')

items=[
('1  INT0 ช่วยนับเวลา', 'ทุก 1,000 รอบจำลองของ CPU จะมีสัญญาณจาก Timer ถ้านาฬิกากำลังเดิน โปรแกรมจะเพิ่มเวลา 1 tick ซึ่งแสดงเป็น 0.01 วินาที'),
('2  INT2 ช่วยบันทึก Lap', 'เมื่อกด Lap โปรแกรมจะคัดลอกเวลาปัจจุบันลงหน่วยความจำ และคำนวณระยะเวลาจาก Lap ก่อนหน้า โดยนาฬิกายังอยู่ในสถานะเดินต่อ'),
]
for lead,body in items:
    p=doc.add_paragraph()
    p.add_run(lead+'  ').bold=True
    p.add_run(body)
p=doc.add_paragraph()
p.add_run('ตัวอย่าง ').bold=True
p.add_run('กด Lap ตอน 3.20 วินาที → เก็บค่า 3.20 → เวลายังเดินต่อเป็น 3.21, 3.22, …')
doc.core_properties.title='การทำงานของ Interrupt ในนาฬิกาจับเวลา'
doc.core_properties.subject='คำอธิบายภาษาไทยพร้อมแผนภาพหนึ่งหน้า'
doc.core_properties.author=''
for root in (doc.styles.element,doc.element):
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
